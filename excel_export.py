"""
Excel export for the drawing pipeline.

Takes the rows built by the web app (or, later, Script B) and returns an .xlsx as
bytes. Kept separate from any UI so Script B can reuse it unchanged.

Two sheets, mirroring the pipeline's own rule that a partly-read drawing is never
classified:
    "Extraction Summary"  DEFENSE and NON_DEFENSE drawings, all required fields read.
    "Needs Attention"     NEEDS_REVIEW and FAILED drawings, with the reason.
"""

from __future__ import annotations

import io
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import drawing_common as dc

# Built from dc.FIELD_LABELS rather than listed by hand, so a field added to the
# extraction engine automatically gets a column here too.
_FIELD_COLS = [dc.FIELD_LABELS[f] for f in dc.FIELD_ORDER]
SUMMARY_COLUMNS = (["Filename"] + _FIELD_COLS
                  + ["Classification", "Classification Reason",
                     "Extraction Method", "Manually Edited", "Date Processed"])
ATTENTION_COLUMNS = (["Filename", "Status", "Reason"] + _FIELD_COLS
                    + ["Extraction Method", "Manually Edited", "Date Processed"])

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_DEFENSE_FILL = PatternFill("solid", fgColor="F8CBAD")


def _write_sheet(ws, columns: List[str], rows: List[Dict[str, str]]) -> None:
    ws.append(columns)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append([row.get(c, "") for c in columns])

    for idx, col in enumerate(columns, start=1):
        longest = max([len(str(col))] + [len(str(r.get(col, ""))) for r in rows])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 12), 60)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def build_workbook(rows: List[Dict[str, str]]) -> bytes:
    """rows: one dict per drawing, with a "Status" key (routing folder name)."""
    good = [r for r in rows if r["Status"] in ("DEFENSE", "NON_DEFENSE")]
    bad = [r for r in rows if r["Status"] not in ("DEFENSE", "NON_DEFENSE")]

    wb = Workbook()
    ws = wb.active
    ws.title = "Extraction Summary"
    good_out = [dict(r, Classification=r["Status"]) for r in good]
    _write_sheet(ws, SUMMARY_COLUMNS, good_out)
    cls_col = SUMMARY_COLUMNS.index("Classification") + 1
    for r_idx in range(2, ws.max_row + 1):
        if ws.cell(r_idx, cls_col).value == "DEFENSE":
            ws.cell(r_idx, cls_col).fill = _DEFENSE_FILL

    _write_sheet(wb.create_sheet("Needs Attention"), ATTENTION_COLUMNS, bad)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
